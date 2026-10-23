from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pathlib import Path
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'docs/presentation';OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
C={'bg':'F4F9F5','ink':'173A35','sub':'58726A','mint':'CFE5D5','green':'B8D6B2','aqua':'A6D2C8','sea':'8FC6B7','teal':'6EA89E','white':'FFFFFF','line':'D7E7DD','code':'153B35','yellow':'FFD166'}
FONT='Noto Sans CJK KR';MONO='DejaVu Sans Mono'
def rgb(c):return RGBColor.from_string(C.get(c,c))
def shape(sl,x,y,w,h,fill='white',line=None,round=True):
 a=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));a.fill.solid();a.fill.fore_color.rgb=rgb(fill);a.line.fill.background() if not line else None
 if line:a.line.color.rgb=rgb(line)
 if round:
  try:a.adjustments[0]=.14
  except:pass
 return a
def txt(sl,x,y,w,h,text,size=18,color='ink',bold=False,font=FONT,align=None):
 a=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=a.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(7)
  if align is not None:p.alignment=align
 return a
def base(title,kicker='PRODUCT & ENGINEERING',sub=None):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb('bg')
 shape(sl,.48,.44,.08,.31,'teal',round=False);txt(sl,.72,.42,11,.3,kicker,10,'sub',True)
 txt(sl,.65,.98,12.1,.7,title,29,'ink',True)
 if sub:txt(sl,.68,1.76,12,.5,sub,13,'sub')
 sl.notes_slide.notes_text_frame.text=f'말줍 프로젝트 정리 | 코드 기준 b65688c | 2026-10-23 (Asia/Seoul)\n{title}\n'
 return sl
def note(sl,text):sl.notes_slide.notes_text_frame.text+='\n'+text

def card(sl,x,y,w,h,tag,title,body,fill='white'):
 shape(sl,x,y,w,h,fill,'line');txt(sl,x+.23,y+.2,w-.46,.3,tag,10,'sub',True);txt(sl,x+.23,y+.72,w-.46,.75,title,20,'ink',True);txt(sl,x+.23,y+1.55,w-.46,h-1.68,body,15,'sub')
def band(sl,text,y=6.38):shape(sl,.68,y,11.98,.48,'mint');txt(sl,.88,y+.09,11.55,.3,text,12,'ink')
def bullets(sl,x,y,w,items,size=18,gap=.7):
 for i,(head,body) in enumerate(items):
  yy=y+i*gap;shape(sl,x,yy+.08,.075,.19,'teal',round=False);txt(sl,x+.22,yy,w-.22,.32,head,size,'ink',True)
  if body:txt(sl,x+.22,yy+.39,w-.22,gap-.42,body,size-3,'sub')
def pic(sl,name,x,y,w,h):
 path=OUT/'screenshots'/name
 iw,ih=Image.open(path).size;r=min(w/iw,h/ih);nw,nh=iw*r,ih*r
 sl.shapes.add_picture(str(path),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))
def table(sl,headers,rows,widths=None,x=.68,y=2.25,w=11.98,rowh=.65):
 widths=widths or [w/len(headers)]*len(headers);xx=x
 for h,ww in zip(headers,widths):shape(sl,xx,y,ww,.52,'ink',round=False);txt(sl,xx+.13,y+.1,ww-.26,.32,h,12,'white',True);xx+=ww
 for i,row in enumerate(rows):
  xx=x;yy=y+.52+i*rowh
  for value,ww in zip(row,widths):shape(sl,xx,yy,ww,rowh,'white' if i%2==0 else 'mint',round=False);txt(sl,xx+.13,yy+.12,ww-.26,rowh-.13,value,12,'ink');xx+=ww

def code(sl,title,content,path,explain):
 shape(sl,.68,2.26,7.8,3.98,'code');txt(sl,.96,2.48,7.25,.3,title,12,'aqua',True,font=MONO)
 txt(sl,.96,3.02,7.25,3.05,content,12,'white',font=MONO)
 txt(sl,8.85,2.4,3.7,.55,'구현 포인트',19,'ink',True);txt(sl,8.85,3.15,3.65,2.9,explain,16,'sub')
 band(sl,'코드 위치: '+path+'  ·  실제 핵심 구문을 축약·줄바꿈하여 표시')
 note(sl,'원본: https://github.com/Min8309/esl_101/blob/main/'+path+'\n'+content+'\n'+explain)
# 01
s=base('좋아하는 한마디를 주워,\n내 말로.','MALJUP · PROJECT DOCUMENTATION');txt(s,.7,3.0,7.3,.85,'말줍 프로젝트 기획·제작 보고서',25,'ink',True);txt(s,.73,4.14,7.2,1.15,'기획 의도  /  사이트맵  /  UX·디자인\n기능 명세  /  핵심 코드  /  배포·운영',19,'sub')
s.shapes.add_picture(str(ROOT/'assets/maljup-logo.png'),Inches(9.45),Inches(2.1),width=Inches(2.3))
band(s,'2026.10.23  ·  구현 코드 기준 b65688c  ·  https://esl-101.onrender.com/')
note(s,'본 자료는 현재 저장소 코드와 로컬 실행 화면을 기준으로 작성했습니다. 실제 Render 환경 변수·유료 자원 설정, 실물 iPhone 마이크·푸시 수신 여부는 별도 검증 대상입니다. 예시 화면은 데모 데이터입니다.')
# 02
s=base('말줍은 영상 속 표현을 학습 루틴으로 바꿉니다.','01 / 기획 개요', '공개 자막에서 발견한 표현을 수집하고, 메모·발화·복습으로 연결하는 개인 학습 웹앱')
card(s,.68,2.45,3.83,3.5,'DISCOVER','스크립트로 발견','유튜브 주소 입력\n시간별 자막 읽기·검색\n관심 단어·문장 선택')
card(s,4.75,2.45,3.83,3.5,'COLLECT','내 표현으로 수집','5색 형광펜·카테고리\n문장 카드·학습 메모\n원본 영상·시각 연결')
card(s,8.82,2.45,3.83,3.5,'PRACTICE','말하고 다시 만나기','직접 녹음·미리듣기\n다음 3일 복습\n시트 내보내기·문장 알림')
#03
s=base('“본 표현”이 “쓸 수 있는 표현”이 되도록','01 / 기획 의도')
table(s,['학습 과정의 불편','제품의 대응','기대 효과'],[['영상과 메모 앱을 반복해서 오감','스크립트·선택·수집을 한 화면에','표현 발견에서 저장까지 동선 단축'],['문장만 저장해 상황을 잊음','영상 URL·타임스탬프 함께 보관','원래 발화 맥락으로 다시 이동'],['단어를 모으고 복습하지 않음','오늘의 복습·체크리스트·달력','수집 이후의 학습 행동 제공'],['내 발화를 기록하기 번거로움','카드 녹음 + 보관함에서 자유 녹음','발화 연습과 파일 관리 연결']],[3.7,4.1,4.18],rowh=.83)
band(s,'핵심 원칙: 발견 → 선택 → 수집 → 발화 → 복습을 짧고 반복 가능한 흐름으로 만든다.')
#04
s=base('PC에서는 깊게 읽고, 아이폰에서는 가볍게 이어갑니다.','01 / 사용자·활용 시나리오')
card(s,.68,2.35,5.88,3.8,'DESKTOP','영상으로 영어를 공부하는 사용자','주소로 스크립트 불러오기\n문장 검색·선택·색상 분류\n카드에 뜻·문맥 메모 남기기\n선택한 문장에 휴대폰 알림 예약')
card(s,6.79,2.35,5.87,3.8,'IPHONE / MOBILE','이동 중 표현을 복습하는 사용자','홈 화면 아이콘으로 앱 열기\n작은 메모판에서 문장 수집·삭제\n짧게 녹음하고 저장 파일 재생\n예약 알림을 눌러 해당 문장 열기')
band(s,'현재 전체 자료는 기기별 저장소에 보관합니다. 알림으로 예약한 문장 사본만 연결 기기로 전달됩니다.')
#05
s=base('구현된 범위와 앞으로의 확장을 분리했습니다.','01 / 범위 정의')
table(s,['구분','현재 범위'],[['구현 완료','카드·메모·3일 복습·달력·체크리스트·녹음·모바일 수집판'],['외부 서비스 설정 필요','Supadata 자막 API / Google Sheets OAuth·API 설정'],['코드 구현 + 운영 설정 필요','웹 푸시: VAPID·Render Disk·상시 실행·아이폰 권한'],['현재 제공하지 않음','영상 화면 캡처, 음성 AI 받아쓰기, AI 해설·발음 채점'],['자동 동기화 아님','Sheets는 새 시트 수동 내보내기 / 전체 기기 동기화 미제공']],[3.0,8.98],rowh=.68)
note(s,'초기 아이디어의 캡처+AI 받아쓰기 통합은 현재 구현되지 않았습니다. 공개 자막과 SRT/VTT 파일을 읽는 기능이 구현 범위입니다. 푸시가 구현되었다는 사실과 실제 Render/아이폰에서 작동 검증이 끝났다는 사실을 혼동하지 않습니다.')
#06
s=base('하나의 문장이 쌓이고, 학습 행동으로 이어지는 흐름','02 / 핵심 사용자 여정')
steps=[('01','주소 입력'),('02','스크립트 탐색'),('03','형광펜 선택'),('04','문장 수집'),('05','메모·녹음'),('06','복습·알림')]
for i,(n,t) in enumerate(steps):
 x=.7+i*2.06;shape(s,x,2.75,1.84,1.5,'white','line');txt(s,x+.18,2.95,1.5,.3,n,12,'teal',True);txt(s,x+.18,3.48,1.52,.5,t,16,'ink',True)
 if i<5:txt(s,x+1.87,3.3,.22,.4,'›',22,'teal')
bullets(s,.82,4.95,11.6,[('수집의 기준','같은 문장·영상·시각·카테고리는 중복 저장하지 않고 기존 자료로 안내합니다.'),('학습의 기준','복습은 실제 수행 날짜를 기록하며, 알림은 사용자가 선택한 시각에 한 번씩 예약합니다.')],17,.8)
#07
s=base('사이트맵: 7개 메뉴와 스크립트 작업 화면','02 / 정보 구조')
shape(s,.7,2.25,2.55,3.7,'ink');txt(s,.98,2.55,2,.5,'말줍',25,'white',True);txt(s,.98,3.45,2,1.7,'PC: 왼쪽 메뉴\n모바일: 하단 5개 아이콘\n우측 상단: 전체 메뉴',16,'white')
labels=[('학습 대시보드','주소 입력 / 최근 카드 / 통계'),('내 표현 수집함','카드·검색·분류·수정·삭제'),('오늘의 복습','빈칸 확인 / 복습 완료'),('녹음 보관함','바로 녹음 / 저장 파일 목록'),('학습 달력','학습·완료·예정·밀린 복습'),('암기 체크리스트','꼭 외울 문장 / 암기 상태'),('데이터 관리','Sheets / 백업·복원 / 푸시')]
for i,(a,b) in enumerate(labels):
 y=2.24+i*.53;shape(s,3.65,y,8.98,.44,'white','line');txt(s,3.88,y+.06,2.7,.3,a,13,'ink',True);txt(s,6.7,y+.06,5.6,.3,b,12,'sub')
band(s,'대시보드 → 스크립트 작업 화면 → 문장 카드 상세(뜻·메모·녹음·알림)로 연결')
#08
s=base('메인 화면은 주소 입력과 최근 학습에 집중합니다.','03 / 주요 화면 · 대시보드','현재 로컬 실행 화면 · 설명용 데모 데이터')
pic(s,'desktop-home.png',.7,2.24,8.1,4.55);bullets(s,9.15,2.5,3.4,[('상단 압축','중복 소개 영역을 줄인 전면 헤더'),('주소 기반 시작','언어 선택·보관 스크립트 재열기'),('최근 표현 카드','카드 열기·우측 삭제 버튼')],16,1.15)
#09
s=base('모바일은 스크립트를 읽으며 오른쪽 메모판으로 수집합니다.','03 / 주요 화면 · 모바일')
pic(s,'mobile-memo.png',.85,2.12,3.25,4.85);bullets(s,4.62,2.45,7.65,[('작은 수집판','형광펜·메모 버튼을 눌러 오른쪽 아래에서 열기'),('원형 형광펜 5색','제목 아래에 배치 / 선택 색상과 카테고리 연결'),('선택·드래그·수집','선택 내용 미리보기와 터치 드래그 / 수집 버튼 대안'),('문장 삭제와 개수 안내','수집 문장 선택 후 삭제 / 신규·중복 수집 안내 분리')],17,.91)
note(s,'최근 요청 반영: 별도 메모 타이핑 칸 및 하단 상세 목록은 모바일 수집판에서 제거했습니다. 학습 메모는 문장 카드 상세에서 작성합니다. 실물 iPhone의 선택·드래그는 추가 확인 대상이며 터치 Chromium 회귀 테스트를 실행했습니다.')
#10
s=base('자막은 “공개 자막 가져오기”를 기본으로 합니다.','03 / 기능 · 스크립트')
card(s,.68,2.35,3.83,3.65,'INPUT','주소·파일','YouTube watch / short / live URL 검증\n영어·한국어·일본어 선택\nSRT·VTT 파일 대체 가져오기')
card(s,4.75,2.35,3.83,3.65,'PROVIDER','서버 자막 처리','Supadata: native 모드\nyt-dlp: 로컬 추출기\n시간 단위를 초로 정규화')
card(s,8.82,2.35,3.83,3.65,'READ','읽기·검색','시간별 문장 리스트\n검색·표시한 줄 필터\n원본 영상의 해당 시각으로 이동')
band(s,'Render의 YouTube 봇 확인은 로컬 추출을 막을 수 있습니다. API도 모든 영상의 성공을 보장하지 않습니다.')
#11
s=base('형광펜 색을 정보의 분류 체계로 사용합니다.','03 / 기능 · 선택·분류')
cols=[('FFD166','단어'),('CCEAC3','문장'),('E0F2FE','표현'),('FFD5E2','질문'),('E6D5FF','문법')]
for i,(co,label) in enumerate(cols):
 x=.85+i*2.5;shape(s,x,2.55,2.23,1.83,'white','line');a=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.77),Inches(2.85),Inches(.65),Inches(.65));a.fill.solid();a.fill.fore_color.rgb=rgb(co);a.line.color.rgb=rgb('line');txt(s,x+.22,3.66,1.8,.45,label,19,'ink',True,align=PP_ALIGN.CENTER)
bullets(s,.88,4.92,11.5,[('정확한 부분 선택','cue 번호 + 시작/끝 문자 위치 + 선택 문자열로 저장해 여러 줄도 추적합니다.'),('수집 품질 관리','카테고리와 원본 시각을 함께 저장하며, 중복 입력은 “이미 수집한 문장”으로 알립니다.')],17,.85)
note(s,'기본 카테고리는 단어·문장·표현·질문·문법이며 데스크탑 상세 도구에서 색상별 분류명을 설정할 수 있습니다. 모바일 간단 수집판은 불필요한 설정을 숨깁니다.')
#12
s=base('학습 카드에는 문장뿐 아니라 맥락을 남깁니다.','03 / 기능 · 수집함·메모')
pic(s,'desktop-script.png',.7,2.25,8.05,4.5);bullets(s,9.1,2.4,3.45,[('카드 기본 정보','표현·뜻·문맥 메모·출처'),('맥락 복원','영상 링크·시작 시간·색상 분류'),('학습 이력','날짜별 암기 메모·녹음·복습'),('편집·정리','검색·필터·개별/여러 개 삭제')],15,.94)
#13
s=base('세 번의 복습과 암기 상태를 함께 관리합니다.','03 / 기능 · 복습·달력·체크리스트')
for i,(a,b) in enumerate([('첫 학습','카드 수집'),('1일 복습','빈칸 확인'),('2일 복습','다시 발화'),('3일 복습','완료 기록')]):
 x=.75+i*3.15;shape(s,x,2.55,2.86,1.48,'mint' if i==0 else 'white','line');txt(s,x+.22,2.85,2.4,.4,a,21,'ink',True);txt(s,x+.22,3.45,2.4,.3,b,14,'sub')
bullets(s,.9,4.64,11.4,[('날짜 중심 복습','다음 날짜부터 복습 가능 / 하루에 중복 완료하지 않도록 검증'),('달력과 암기 체크리스트','실제 복습·예정·밀린 일정 표시 / 꼭 외울 문장과 암기 완료 상태는 별도 관리')],17,.85)
band(s,'“매일 다음 3일 복습”과 “1·3·7일 간격 반복”은 다릅니다. 1·3·7일 알림은 각각 수동 예약합니다.')
#14
s=base('녹음 보관함에서 바로 말하고, 아래 목록에서 다시 듣습니다.','03 / 기능 · 녹음')
pic(s,'desktop-recordings.png',.7,2.2,7.35,4.64);pic(s,'mobile-recordings.png',8.3,2.12,2.17,4.7);txt(s,10.8,2.5,1.82,3.5,'자유 녹음\n카드 연결\n마이크 권한\n미리듣기\n명시적 저장\n다운로드·삭제',13,'sub')
note(s,'MediaRecorder 기반으로 녹음하며 iOS에서는 지원되는 audio/mp4를 우선 선택합니다. 저장은 IndexedDB이며 녹음 보관함은 최신순입니다. 카드 상세의 따라 말하기 녹음도 유지합니다. 화면의 데모 파일은 생성한 무음 WAV로 실제 마이크 녹음이 아닙니다. 아이폰 실물 마이크·오디오 재생은 별도 검증이 필요합니다.')
#15
s=base('Google Sheets는 문장·메모를 내 계정에 내보내는 경로입니다.','03 / 기능 · Google Sheets')
for i,(a,b) in enumerate([('Google 계정 연결','OAuth drive.file 권한'),('저장 버튼 실행','현재 카드 전체 변환'),('새 스프레드시트','16개 열로 기록'),('생성된 시트 열기','표로 확인·편집')]):
 x=.7+i*3.15;card(s,x,2.45,2.91,2.65,str(i+1).zfill(2),a,b)
txt(s,.85,5.45,11.7,.65,'표현·뜻·학습 메모·날짜별 메모·영상 정보·복습·암기 상태·분류·색상 포함',17,'ink',True)
band(s,'수동 내보내기 / 저장마다 새 시트 / 녹음 제외 / 시트 수정 내용은 앱으로 역동기화하지 않음')
note(s,'클라이언트 공개 ID와 승인된 HTTPS 원본이 필요하며 Google Sheets API를 활성화해야 합니다. 사용자는 연결 성공을 보고했으나 본 문서는 사용자 계정의 파일 내용이나 현재 권한을 독립 검증한 자료가 아닙니다.')
#16
s=base('예약한 문장을 연결된 아이폰의 웹 푸시로 전달합니다.','03 / 기능 · 문장 알림','코드 구현 완료 · 실제 운영에는 Render·기기 설정과 수신 검증 필요')
card(s,.7,2.45,3.82,3.5,'PAIR','기기 연결 키','아이폰에서 연결 키 생성\n알림 구독 등록·테스트\nPC에 같은 키 입력')
card(s,4.76,2.45,3.82,3.5,'SCHEDULE','한 번씩 날짜·시간 예약','카드 문장·뜻·메모 사본\nUTC 예약 시각 저장\n목록 확인·예약 취소')
card(s,8.82,2.45,3.82,3.5,'DELIVER','알림 → 해당 문장','서버가 30초마다 확인\nVAPID·Web Push 전송\n누르면 문장 열기·수집함 저장')
band(s,'SQLite 영구 Disk + 상시 실행 서버 필요 / 집중 모드·네트워크에 따라 표시·수신이 지연될 수 있음')
#17
s=base('아이폰 11: Safari에서 홈 화면에 추가한 뒤 알림을 허용합니다.','03 / 설치·온보딩')
items=[('01','iOS 16.4 이상 확인','설정 → 일반 → 정보'),('02','Safari에서 말줍 열기','esl-101.onrender.com'),('03','공유 → 홈 화면에 추가','말줍 이름으로 추가'),('04','홈 화면 아이콘으로 실행','데이터 관리 → 아이폰 문장 알림'),('05','연결 키 생성·알림 허용','테스트 알림으로 실제 수신 확인')]
for i,(n,a,b) in enumerate(items):
 y=2.2+i*.79;shape(s,.75,y,.5,.5,'mint');txt(s,.85,y+.1,.3,.25,n,11,'ink',True);txt(s,1.58,y+.03,5,.4,a,19,'ink',True);txt(s,7.0,y+.07,5.5,.35,b,15,'sub')
note(s,'Render 사전 설정: VAPID_PUBLIC_KEY, VAPID_PRIVATE_KEY, VAPID_SUBJECT, PUSH_DATA_DIR=/var/data 및 영구 Disk, 유휴 정지되지 않는 인스턴스. 비용은 요금제를 확인하며 앱 기본 알람 같은 강제 소리는 아닙니다.')
#18
s=base('부드러운 민트와 명확한 인터랙션으로 통일합니다.','04 / 디자인 시스템')
for i,(name,hexx) in enumerate([('Soft Mint','CFE5D5'),('Pastel Green','B8D6B2'),('Aqua Breeze','A6D2C8'),('Seafoam','8FC6B7'),('Teal Mist','6EA89E')]):
 x=.72+i*2.5;shape(s,x,2.45,2.26,1.35,hexx);txt(s,x,4.02,2.3,.38,name,16,'ink',True);txt(s,x,4.53,2.3,.3,'#'+hexx,13,'sub',font=MONO)
bullets(s,.85,5.3,11.5,[('컴포넌트 원칙','둥근 카드·명료한 버튼 / 모바일 원형 형광펜 / 채워진 마이크 아이콘'),('내비게이션 원칙','PC 왼쪽 메뉴 유지 / 모바일 하단 아이콘 + 상단 전체 메뉴 / 스크립트 화면에 학습 공간 집중')],16,.72)
#19
s=base('브라우저 학습 데이터와 서버 연동 기능을 분리합니다.','05 / 시스템 아키텍처')
card(s,.7,2.35,3.65,3.8,'CLIENT','브라우저','HTML·CSS·ES Modules\nIndexedDB: 카드·스크립트·녹음\nService Worker: 알림 수신\nlocalStorage: 공개 설정·연결 키')
card(s,4.82,2.35,3.65,3.8,'SERVER','Node.js 24','정적 파일·HTTP API\n자막 제공자 선택·검증\nWeb Push 예약 처리\nSQLite: 구독·예약')
card(s,8.94,2.35,3.65,3.8,'SERVICES','외부 서비스','Supadata / YouTube 자막\nGoogle Identity·Sheets\nAPNs 등 웹 푸시 제공자\nRender Docker·영구 Disk')
band(s,'Google Sheets 저장은 브라우저 → Google 직접 호출 / 자막 API 키·VAPID 비밀 키는 서버 환경 변수')
#20
s=base('데이터 모델: 개인 학습 원본은 IndexedDB에 저장합니다.','05 / 데이터 구조')
table(s,['저장 위치','엔터티','대표 필드'],[['IndexedDB / cards','표현 카드','id, text, meaning, note, url, start, category, reviews, next'],['IndexedDB / recordings','오디오 파일','id, cardId, title, blob, created'],['IndexedDB / scripts','자막·형광펜','id, data.cues, highlights, updated / palette'],['SQLite / subscriptions','푸시 구독','endpoint, owner(키 해시), data, created'],['SQLite / reminders','예약 알림','id, owner, card, due, status, attempts, updated']],[2.5,2.25,7.23],rowh=.7)
band(s,'IndexedDB 버전 2 / 구버전 자료 유지 / 전체 JSON 백업은 카드·복습·녹음·스크립트·형광펜 포함')
#21
s=base('제작 방법: 작은 모듈을 연결하고 실제 동작을 검증합니다.','05 / 제작 방식')
table(s,['단계','제작 방법','산출물'],[['구조 설계','화면 흐름·저장 구조·접근 권한 구분','사이트맵 / 데이터 모델'],['기능 개발','Vanilla JS UI + Node HTTP API + Python 추출기','src/ · server.mjs · scripts/'],['품질 확인','순수 로직·HTTP·영구 저장·브라우저 회귀 검사','Node / Python / Chromium 테스트'],['배포','잠금 파일 기반 설치·Docker·비관리자 실행','Dockerfile / Render 안내'],['운영 문서화','키 설정·아이폰 설치·한계와 복구 경로 기록','README / 배포·푸시 설정 문서']],[1.7,6.3,3.98],rowh=.7)
#22
s=base('핵심 코드 01 — 자막 입력 검증과 제공자 선택','06 / 핵심 구현')
code(s,'youtube-transcript.mjs',"const input = transcriptInput(value, language);\nconst provider = process.env.TRANSCRIPT_PROVIDER\n  || (process.env.SUPADATA_API_KEY?.trim()\n      ? 'supadata' : 'local');\nif (provider === 'supadata')\n  return supadata(input);",'youtube-transcript.mjs','YouTube 영상 ID·언어를 먼저 검증합니다.\n\nAPI 키가 있으면 Supadata, 명시한 local 설정이면 Python 추출기를 사용합니다.\n\n잘못된 설정은 명확한 오류로 처리합니다.')
#23
s=base('핵심 코드 02 — 정확한 부분 선택과 중복 수집 방지','06 / 핵심 구현')
code(s,'selection anchor / duplicate guard',"// highlighted substring anchor\n{ cue, start, end, text }\n\n// same phrase, video, time and category\nc.text === item.text\n  && sameVideo(c.url, item.url)\n  && c.start === item.start\n  && (c.category || '일반') === item.category",'src/transcript-workspace.js · src/app.js','화면 위치 대신 문장 번호·문자 오프셋으로 저장합니다.\n\n원본과 선택 문자열이 일치해야 유효합니다.\n\n같은 자료를 다시 수집하면 0개 성공 대신 중복 안내를 표시합니다.')
#24
s=base('핵심 코드 03 — 녹음 상태와 늦은 권한 응답 처리','06 / 핵심 구현')
code(s,'ArchiveRecorder.start()',"const stream = await navigator.mediaDevices\n  .getUserMedia({ audio: true });\nif (session !== this.session\n    || !root.isConnected) {\n  stream.getTracks().forEach(t => t.stop());\n  return;\n}\nconst recorder = new MediaRecorder(stream, options);",'src/archive-recorder.js','idle → 권한 요청 → recording → preview → saving\n\n권한 요청 중 화면을 벗어나면 늦게 열린 마이크도 즉시 해제합니다.\n\nMediaRecorder는 원본에서 지원 MIME을 고른 뒤 options를 구성합니다.')
#25
s=base('핵심 코드 04 — Google Sheets에 안전한 값으로 쓰기','06 / 핵심 구현')
code(s,'Sheets API / updateCells',"fetch('https://sheets.googleapis.com/v4/'\n  + 'spreadsheets' + path, {\n  method: 'POST',\n  headers: { Authorization: `Bearer ${token}` },\n  body: JSON.stringify(body)\n});\n// Cells use stringValue / numberValue,\n// not interpreted formulas.",'src/google-sheets.js','OAuth drive.file 범위를 사용합니다.\n\n새 시트를 만든 뒤 16개 열의 값으로 기록합니다.\n\n문자열은 수식으로 실행하지 않으며 실패하면 부분 생성된 시트도 안내합니다.')
#26
s=base('핵심 코드 05 — 기기 연결 키와 예약 상태를 분리','06 / 핵심 구현')
code(s,'push-service.mjs',"createHash('sha256')\n  .update(token).digest('hex');\n\nSELECT * FROM reminders\n WHERE status = 'pending' AND due <= ?\n ORDER BY due LIMIT 20;\n\n// pending → sending → sent / failed / expired",'push-service.mjs · sw.js','연결 키 원문은 서버 DB에 저장하지 않습니다.\n\n조회·삭제에도 같은 키의 해시를 적용합니다.\n\n30초 확인, 만료 구독 삭제, 제한된 재시도로 상태를 관리합니다.')
#27
s=base('배포: GitHub main → Render Docker Web Service','07 / 배포·환경 설정')
table(s,['영역','필수 설정','확인 방법'],[['기본 실행','Node 24 / npm ci / Docker / PORT','/healthz → status: ok'],['자막 API','TRANSCRIPT_PROVIDER / SUPADATA_API_KEY','공개 자막 영상·선택 언어로 조회'],['Google Sheets','웹 OAuth ID / HTTPS 원본 / Sheets API 활성화','계정 연결 → 새 시트 저장·열기'],['문장 알림','VAPID 3개 값 / PUSH_DATA_DIR=/var/data','/api/push/config → enabled: true'],['예약 유지','상시 실행 서비스 / 영구 Disk / 단일 인스턴스','재시작 후 예약·실제 아이폰 수신 확인']],[2.0,5.5,4.48],rowh=.7)
band(s,'GitHub 푸시 완료와 Render 배포 Live, 실제 외부 연동 성공은 각각 확인해야 합니다.')
#28
s=base('보안·개인정보: 저장 위치와 권한 범위를 보이게 합니다.','07 / 보안·권한')
card(s,.68,2.4,3.83,3.58,'GOOGLE','최소 파일 권한','drive.file 범위\n비밀번호를 앱에 저장하지 않음\n접근 토큰은 브라우저 메모리\n클라이언트 ID는 공개 식별자')
card(s,4.75,2.4,3.83,3.58,'SERVER','비밀 키와 입력 검증','Supadata·VAPID 키는 환경 변수\n정적 파일 경로 허용 목록\n자막·푸시 요청 크기·횟수 제한\n푸시 대상 HTTPS 도메인 검증')
card(s,8.82,2.4,3.83,3.58,'DATA','사용자가 통제','녹음은 브라우저에만 저장\n예약한 문장 사본은 서버 저장\n푸시 연결 키는 비밀번호처럼 관리\n서버 DB 앱 수준 암호화는 없음')
note(s,'전문 보안 감사를 완료했다는 의미는 아닙니다. 공개 앱의 로그인·다중 사용자 권한은 향후 강화 대상입니다. Web Push 전송은 표준 암호화를 사용하며 서버 저장 자료는 운영자/호스팅 계정 접근 통제가 필요합니다. 외부 API 비용·할당량 제한도 별도로 설정해야 합니다.')
#29
s=base('검증: 저장·권한·모바일 회귀를 함께 확인합니다.','07 / 테스트·품질')
for x,n,label in [(.75,'41','Node 자동 테스트'),(4.95,'6','Python 자동 테스트'),(9.13,'5','반응형 화면 폭')]:
 shape(s,x,2.4,3.5,1.62,'white','line');txt(s,x+.25,2.57,2.9,.73,n,38,'teal',True);txt(s,x+.25,3.45,3,.3,label,14,'sub')
bullets(s,.9,4.5,11.5,[('Chromium 브라우저 검사','카드 삭제 취소/확인, 모바일 터치 수집·삭제, 중복 안내, 녹음 보관함 저장·재접속·삭제'),('검증의 범위','마이크·Google·자막 API는 모의 응답을 포함합니다. 실물 iPhone과 실제 계정·푸시 수신은 별도 확인합니다.')],16,.87)
note(s,'검증 스크립트: npm run check / npm test / python3 -m unittest discover -s tests -p test_*.py / npm run test:browser. 반응형 폭 1440·1024·768·390·320px. 숫자는 코드 기준 시점의 실행 결과이며 향후 변경될 수 있습니다.')
#30
s=base('운영에서 남은 확인과 다음 확장 우선순위','07 / 한계·개선 방향')
table(s,['우선순위','내용','현재 상태'],[['P0 · 운영 검증','Render 최신 배포 / Sheets 실데이터 / 자막 제공자 응답','실제 계정·배포 환경에서 확인 필요'],['P0 · 아이폰 검증','마이크 녹음·재생 / 홈 화면 설치 / 푸시 실수신','브라우저 모의 검사와 별도'],['P1 · 데이터 연결','로그인·사용자별 클라우드 DB·기기 간 동기화','현재 전체 자료는 브라우저별 분리'],['P1 · 운영 안정성','공유 DB·전용 작업자·모니터링·백업','현재 SQLite 단일 인스턴스 기준'],['P2 · 학습 확장','간격 반복·AI 해설·발음 피드백·캡처','현재 미구현 / 별도 기획·비용 검토']],[2.1,6.3,3.58],rowh=.7)
#31
s=base('소스 코드와 설정 문서 찾아보기','08 / 개발·운영 인수인계')
table(s,['파일','담당 기능'],[['src/app.js / src/core.js','화면·IndexedDB·카드 / 날짜·복습·URL·자막 파싱'],['src/transcript-workspace.js','부분 선택·형광펜·드래그 수집·모바일 메모판'],['src/archive-recorder.js / src/recording-ui.js','보관함 직접 녹음·상태 관리 / 마이크 아이콘'],['src/google-sheets.js / youtube-transcript.mjs','OAuth·Sheets 내보내기 / 자막 제공자 선택'],['push-service.mjs / src/push-notifications.js / sw.js','예약 DB·전송 / 기기 연결 UI / 알림 수신·열기'],['docs/render-deploy.md / docs/iphone-push.md','Render 배포·자막 API / 아이폰 설치·푸시 설정']],[6.0,5.98],rowh=.61)
note(s,'저장소: https://github.com/Min8309/esl_101\n서비스: https://esl-101.onrender.com/\n각 파일의 최신 코드는 저장소에서 확인합니다. 발표 자료는 기준 커밋의 스냅샷입니다.')
#32
s=base('발표·시연 순서','08 / 데모 가이드')
for i,(a,b) in enumerate([('주소에서 시작','공개 자막이 있는 영상 URL과 언어 선택'),('하나의 표현 수집','원형 형광펜 → 선택 → 드래그 또는 수집 버튼'),('내 표현으로 만들기','내 표현 수집함 → 카드 → 뜻·학습 메모 작성'),('직접 발화하기','녹음 보관함 → 시작·종료·미리듣기·저장'),('자료 보관·다시 만나기','Sheets 내보내기 / 3일 복습 / 설정 완료 시 예약 알림')]):
 y=2.2+i*.8;txt(s,.8,y,.65,.45,str(i+1).zfill(2),23,'teal',True);txt(s,1.75,y+.03,3.65,.4,a,19,'ink',True);txt(s,5.55,y+.07,7.0,.4,b,15,'sub')
# Footer and source notes
for i,s in enumerate(prs.slides,1):
 shape(s,.68,7.06,11.98,.012,'line',round=False);txt(s,.7,7.18,10,.2,'말줍  |  프로젝트 기획·제작 문서  |  2026.10.23',9,'sub');txt(s,11.75,7.13,.9,.24,f'{i:02d} / {len(prs.slides):02d}',10,'sub',align=PP_ALIGN.RIGHT)
prs.core_properties.title='말줍 프로젝트 기획·제작 보고서';prs.core_properties.subject='기획 의도, 제작 방법, 사이트맵, 기능과 핵심 코드';prs.core_properties.author='말줍 프로젝트';prs.core_properties.keywords='말줍,기획,제작,사이트맵,웹앱,영어학습,코드';prs.save(OUT/'maljup-project.pptx')
print(f'Created {len(prs.slides)} slides: {OUT}/maljup-project.pptx')
